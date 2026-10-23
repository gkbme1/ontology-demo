"""工艺限值：上下限 / 四档报警限。

数据边界（与规则、对象映射同一条线）：
- 本体**不生成**任何限值，只加载独立数据源 `data/source/6APA_工艺限值.xlsx`；
- 文件不存在时接口返回 `available: false` 与 `emptyReason`，并给出模板列，绝不编造；
- 每条限值都带「限值依据」「置信度」「确认状态」，只有状态为「已确认」的才算正式工艺依据。

绑定方式：限值挂在工艺属性（关系对象）上，通过本体已有的测量锚点自动落到具体 DCS 位号；
也可以直接按位号写特例，位号支持 `*` 通配（如 TI_1407*）。

只用 openpyxl（本体服务已依赖），不新增依赖。
"""
from __future__ import annotations

import fnmatch
import threading
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[2]
LIMIT_FILE = ROOT / "data" / "source" / "6APA_工艺限值.xlsx"
TEMPLATE_FILE = ROOT / "data" / "source" / "6APA_工艺限值.模板.xlsx"
SHEET = "工艺限值"
SOURCE_TYPE = "LIMIT_FILE"

TEMPLATE_COLUMNS = [
    "限值编号 (Limit ID)", "适用对象类型 (Applies To Type)", "适用对象 (Applies To)", "适用工段 (Area)",
    "适用工况 (Operating Mode)", "参数含义 (Parameter)", "单位 (Unit)",
    "低低限 LL", "低限 L", "正常下限 (Normal Low)", "正常上限 (Normal High)", "高限 H", "高高限 HH",
    "限值依据 (Basis)", "依据类型 (Basis Type)", "置信度 (Confidence)", "确认状态 (Status)", "备注 (Note)",
]

STATUS_CONFIRMED = "已确认"

# 越限判定结果
NORMAL, OUT_OF_NORMAL = "NORMAL", "OUT_OF_NORMAL"
LOW, LOW_LOW, HIGH, HIGH_HIGH = "LOW", "LOW_LOW", "HIGH", "HIGH_HIGH"
UNKNOWN = "UNKNOWN"
STATUS_LABEL = {
    NORMAL: "正常", OUT_OF_NORMAL: "超出正常区间（未越报警限）",
    LOW: "低于低限", LOW_LOW: "低于低低限", HIGH: "高于高限", HIGH_HIGH: "高于高高限",
    UNKNOWN: "无法判断（无数值或该档限值未设）",
}

_cache: dict[str, Any] = {"mtime": None, "items": []}
_lock = threading.Lock()


# --------------------------------------------------------------------- 加载


def _num(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _load() -> list[dict[str, Any]]:
    """读 Excel。文件没变就用缓存；文件不存在返回空列表。"""
    if not LIMIT_FILE.is_file():
        return []
    mtime = LIMIT_FILE.stat().st_mtime
    with _lock:
        if _cache["mtime"] == mtime:
            return _cache["items"]
        from openpyxl import load_workbook

        wb = load_workbook(LIMIT_FILE, read_only=True, data_only=True)
        ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
        items: list[dict[str, Any]] = []
        for index, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if row is None or not _text(row[0] if len(row) else ""):
                continue
            cell = lambda i: row[i] if len(row) > i else None  # noqa: E731
            limit = {
                "id": _text(cell(0)),
                "appliesToType": _text(cell(1)),
                "appliesTo": _text(cell(2)),
                "area": _text(cell(3)),
                "areaCode": None,
                "areaId": None,
                "operatingMode": _text(cell(4)),
                "parameter": _text(cell(5)),
                "unit": _text(cell(6)),
                "lowLow": _num(cell(7)),
                "low": _num(cell(8)),
                "normalLow": _num(cell(9)),
                "normalHigh": _num(cell(10)),
                "high": _num(cell(11)),
                "highHigh": _num(cell(12)),
                "basis": _text(cell(13)),
                "basisType": _text(cell(14)),
                "confidence": _text(cell(15)),
                "status": _text(cell(16)),
                "note": _text(cell(17)),
                "sourceType": SOURCE_TYPE,
                "source": {"file": LIMIT_FILE.name, "sheet": ws.title, "row": index},
            }
            limit["hasValues"] = any(
                limit[k] is not None for k in ("lowLow", "low", "normalLow", "normalHigh", "high", "highHigh")
            )
            limit["confirmed"] = limit["status"] == STATUS_CONFIRMED
            items.append(limit)
        wb.close()
        _cache["mtime"] = mtime
        _cache["items"] = items
        return items


def available() -> bool:
    return bool(_load())


def _empty_payload() -> dict[str, Any]:
    return {
        "items": [], "total": 0, "available": False,
        "emptyReason": "当前暂无工艺限值数据源。限值必须来自工艺规程 / 现场确认，本体不会自动生成。"
                       f"把填好的 {LIMIT_FILE.name} 放入 data/source 后重启服务即自动生效。",
        "expectedSource": LIMIT_FILE.name,
        "templateFile": TEMPLATE_FILE.name if TEMPLATE_FILE.is_file() else None,
        "templateColumns": TEMPLATE_COLUMNS,
    }


# ----------------------------------------------------------- 限值 → DCS 位号


def _tags_of_equipment(store, equipment_id: str) -> list[dict[str, Any]]:
    equip = store.dcs.equipment.get(equipment_id) or {}
    return [store.dcs.tags[t] for t in equip.get("tagIds", []) if t in store.dcs.tags]


def _tags_of_parameter(store, parameter_id: str) -> list[dict[str, Any]]:
    param = store.dcs.parameters.get(parameter_id) or {}
    return [store.dcs.tags[t] for t in param.get("tagIds", []) if t in store.dcs.tags]


def resolve_targets(store, limit: dict[str, Any]) -> tuple[list[dict[str, Any]], list[str]]:
    """把一条限值落到具体 DCS 位号上，返回 (位号列表, 说明)。"""
    kind, name = limit["appliesToType"], limit["appliesTo"]
    notes: list[str] = []
    if not name:
        return [], ["限值未填写适用对象"]

    if kind in ("工艺属性", "质量指标", "业务指标", "关系对象", "公用介质", "执行机构", "设备部件", "控制仪表"):
        entity = next((e for e in store.relations.entities.values() if e["name"] == name), None)
        if entity is None:
            return [], [f"关系表中没有名为「{name}」的对象，限值无法落到测点"]
        seen: set[str] = set()
        tags: list[dict[str, Any]] = []
        for match in entity["match"]["matches"]:
            if match["objectType"] == "tag":
                found = [store.dcs.tags[match["id"]]] if match["id"] in store.dcs.tags else []
            elif match["objectType"] == "parameter":
                found = _tags_of_parameter(store, match["id"])
            elif match["objectType"] == "equipment":
                found = _tags_of_equipment(store, match["id"])
            else:
                found = []
            for tag in found:
                if tag["id"] not in seen:
                    seen.add(tag["id"])
                    tags.append(tag)
        if not tags:
            notes.append("该对象在 DCS 点表中未匹配到测点，限值目前只能人工使用，无法自动做越限判断")
        else:
            strategies = set(entity["match"].get("matchStrategies") or [])
            if "RELATION_TAG_BINDING" in strategies:
                notes.append("对象到测点的对应来自工业关系表自带的「关联DCS位号 / 绑定测点」声明")
            if "MANUAL_MAPPING" in strategies:
                notes.append("对象到测点的对应部分来自人工映射表，映射条目「待现场确认」")
        return tags, notes

    if kind == "DCS点":
        pattern = name.upper()
        if any(ch in pattern for ch in "*?"):
            tags = [t for t in store.dcs.tags.values() if fnmatch.fnmatch((t["name"] or "").upper(), pattern)]
            if not tags:
                notes.append(f"位号通配 {name} 没有匹配到任何测点")
            return tags, notes
        tags = [t for t in store.dcs.tags.values() if (t["name"] or "").upper() == pattern]
        if not tags:
            notes.append(f"点表中没有位号 {name}")
        return tags, notes

    if kind == "设备":
        equips = store.dcs.equipment_by_name(name)
        tags = [t for e in equips for t in _tags_of_equipment(store, e["id"])]
        if not tags:
            notes.append(f"点表中没有设备「{name}」或该设备无测点")
        return tags, notes

    if kind in ("参数", "参数语义"):
        params = store.dcs.parameters_by_name(name)
        tags = [t for p in params for t in _tags_of_parameter(store, p["id"])]
        if not tags:
            notes.append(f"点表中没有参数语义「{name}」")
        return tags, notes

    return [], [f"未知的适用对象类型「{kind}」，可选：工艺属性 / 质量指标 / 业务指标 / 设备 / 参数 / DCS点"]


def _target_brief(tag: dict[str, Any]) -> dict[str, Any]:
    return {
        "tagId": tag["id"], "tagName": tag["name"], "description": tag.get("rawDescription"),
        "equipmentId": tag.get("equipmentId"), "equipmentName": tag.get("equipmentName"),
        "parameterId": tag.get("parameterId"), "parameterName": tag.get("parameterName"),
        "areaName": tag.get("areaName"), "unit": tag.get("unit"), "station": tag.get("station"),
    }


def _with_targets(store, limit: dict[str, Any], include: bool, sample: int = 8) -> dict[str, Any]:
    tags, notes = resolve_targets(store, limit)
    out = {**limit, "targetCount": len(tags)}
    if notes:
        out["notes"] = notes
    if include:
        out["targets"] = [_target_brief(t) for t in tags[:sample]]
        if len(tags) > sample:
            out["targetsTruncated"] = len(tags) - sample
    return out


# --------------------------------------------------------------------- 查询


def _matches_filters(limit: dict[str, Any], q: dict) -> bool:
    if q.get("area"):
        wanted = str(q["area"]).strip()
        # 同时接受工段名称、工段编号 S11、以及 area-S11
        candidates = {limit.get("area"), limit.get("areaCode"), limit.get("areaId")}
        if wanted not in candidates:
            return False
    if q.get("appliesTo") and limit["appliesTo"] != q["appliesTo"]:
        return False
    if q.get("appliesToType") and limit["appliesToType"] != q["appliesToType"]:
        return False
    if q.get("status") and limit["status"] != q["status"]:
        return False
    if q.get("confirmed") in ("true", "false") and limit["confirmed"] != (q["confirmed"] == "true"):
        return False
    if q.get("hasValues") in ("true", "false") and limit["hasValues"] != (q["hasValues"] == "true"):
        return False
    keyword = (q.get("keyword") or "").strip()
    if keyword:
        blob = " ".join(str(limit.get(k) or "") for k in
                        ("appliesTo", "parameter", "area", "operatingMode", "basis", "note"))
        if keyword not in blob:
            return False
    return True


def _resolve_limit_area(store, limit: dict[str, Any]) -> None:
    """把限值的「适用工段」对到 v2 的工段编号上；对不上就留空并保留原文。"""
    if limit.get("areaId") or not limit.get("area"):
        return
    aid = store.dcs.resolve_area(limit["area"])
    if aid:
        limit["areaId"] = aid
        limit["areaCode"] = store.dcs.areas[aid]["code"]


def list_limits(store, q: dict) -> dict[str, Any]:
    items = _load()
    if not items:
        return _empty_payload()
    for limit in items:
        _resolve_limit_area(store, limit)
    hits = [limit for limit in items if _matches_filters(limit, q)]
    include = (q.get("includeTargets") or "true").lower() != "false"
    return {
        "items": [_with_targets(store, limit, include) for limit in hits],
        "total": len(hits),
        "available": True,
        "sourceType": SOURCE_TYPE,
        "source": {"file": LIMIT_FILE.name, "sheet": SHEET},
        "statusSummary": _summary(items),
        "facets": {
            "area": sorted({i["area"] for i in items if i["area"]}),
            "areaCode": sorted({i["areaCode"] for i in items if i.get("areaCode")}),
            "appliesToType": sorted({i["appliesToType"] for i in items if i["appliesToType"]}),
            "status": sorted({i["status"] for i in items if i["status"]}),
            "operatingMode": sorted({i["operatingMode"] for i in items if i["operatingMode"]}),
        },
        "notice": "限值来自独立数据源，状态非「已确认」的条目仅供工艺调查核对，不能直接作为报警依据。",
    }


def _summary(items: list[dict[str, Any]]) -> dict[str, int]:
    out: dict[str, int] = {"total": len(items), "withValues": 0, "confirmed": 0}
    for limit in items:
        if limit["hasValues"]:
            out["withValues"] += 1
        if limit["confirmed"]:
            out["confirmed"] += 1
        key = limit["status"] or "未填写"
        out[key] = out.get(key, 0) + 1
    return out


def limits_for_object(store, object_id: str) -> dict[str, Any]:
    """某个本体对象适用哪些限值：按对象解析出的测点与限值覆盖的测点求交集。"""
    items = _load()
    if not items:
        return {**_empty_payload(), "objectId": object_id}
    for limit in items:
        _resolve_limit_area(store, limit)

    scope, object_tags, notes = _object_tags(store, object_id)
    tag_ids = {t["id"] for t in object_tags}
    name = scope.get("objectName") or ""

    hits = []
    for limit in items:
        targets, target_notes = resolve_targets(store, limit)
        covered = [t for t in targets if t["id"] in tag_ids]
        direct = limit["appliesTo"] == name and limit["appliesToType"] in (
            "工艺属性", "质量指标", "业务指标", "关系对象", "设备", "参数", "参数语义")
        if not covered and not direct:
            continue
        entry = {**limit, "targetCount": len(targets), "matchedTargetCount": len(covered),
                 "matchedTargets": [_target_brief(t) for t in covered[:8]]}
        if target_notes:
            entry["notes"] = target_notes
        if not covered and direct:
            entry.setdefault("notes", []).append("该限值直接挂在本对象上，但未落到具体测点")
        hits.append(entry)

    return {
        "scope": scope, "items": hits, "total": len(hits), "available": True,
        "objectTagCount": len(object_tags), "notes": notes,
        "sourceType": SOURCE_TYPE, "source": {"file": LIMIT_FILE.name, "sheet": SHEET},
        "notice": "限值来自独立数据源，状态非「已确认」的条目仅供工艺调查核对，不能直接作为报警依据。",
    }


def _object_tags(store, object_id: str) -> tuple[dict[str, Any], list[dict[str, Any]], list[str]]:
    """复用与实时值一致的对象解析；实时值模块不存在时退化为本地实现。"""
    try:
        from backend.services.timeseries import resolve_object

        return resolve_object(store, object_id)
    except ImportError:  # pragma: no cover - 没装实时值补丁时的兜底
        pass

    dcs = store.dcs
    if object_id.startswith(":"):
        equip = dcs.equipment_by_owl(object_id)
        object_id = equip["id"] if equip else object_id
    if object_id.startswith("tag-") and object_id in dcs.tags:
        tag = dcs.tags[object_id]
        return {"objectId": object_id, "objectKind": "tag", "objectName": tag["name"]}, [tag], []
    if object_id.startswith("param-"):
        return ({"objectId": object_id, "objectKind": "parameter",
                 "objectName": (dcs.parameters.get(object_id) or {}).get("name", "")},
                _tags_of_parameter(store, object_id), [])
    if object_id.startswith("equip-"):
        return ({"objectId": object_id, "objectKind": "equipment",
                 "objectName": (dcs.equipment.get(object_id) or {}).get("name", "")},
                _tags_of_equipment(store, object_id), [])
    if object_id.startswith("robj-"):
        entity = store.relations.entities.get(object_id) or {}
        tags: list[dict[str, Any]] = []
        for match in entity.get("match", {}).get("matches", []):
            tags += _tags_of_parameter(store, match["id"]) if match["objectType"] == "parameter" else \
                _tags_of_equipment(store, match["id"])
        return {"objectId": object_id, "objectKind": "relationEntity", "objectName": entity.get("name", "")}, tags, []
    resolved_area = dcs.resolve_area(object_id)
    if resolved_area:
        area = dcs.areas.get(resolved_area, {})
        return ({"objectId": resolved_area, "objectKind": "area",
                 "objectName": area.get("name") or resolved_area,
                 "areaCode": area.get("code")},
                [t for t in dcs.tags.values() if t.get("areaId") == resolved_area], [])
    return {"objectId": object_id, "objectKind": "unknown", "objectName": ""}, [], ["无法识别的对象 ID"]


# ----------------------------------------------------------------- 越限判断


def evaluate(limit: dict[str, Any], value: Any) -> dict[str, Any]:
    """把一个实测值放到限值里判断。判断不了就说判断不了，不猜。"""
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        return {"status": UNKNOWN, "statusLabel": STATUS_LABEL[UNKNOWN], "reason": "该测点没有可比较的数值"}
    if not limit["hasValues"]:
        return {"status": UNKNOWN, "statusLabel": STATUS_LABEL[UNKNOWN], "reason": "该条限值尚未填写数值"}

    checks = [
        (limit["highHigh"], lambda v, t: v > t, HIGH_HIGH),
        (limit["lowLow"], lambda v, t: v < t, LOW_LOW),
        (limit["high"], lambda v, t: v > t, HIGH),
        (limit["low"], lambda v, t: v < t, LOW),
    ]
    for threshold, test, status in checks:
        if threshold is not None and test(value, threshold):
            return {"status": status, "statusLabel": STATUS_LABEL[status],
                    "threshold": threshold, "deviation": round(value - threshold, 4)}

    low, high = limit["normalLow"], limit["normalHigh"]
    if low is not None and value < low:
        return {"status": OUT_OF_NORMAL, "statusLabel": STATUS_LABEL[OUT_OF_NORMAL],
                "threshold": low, "deviation": round(value - low, 4)}
    if high is not None and value > high:
        return {"status": OUT_OF_NORMAL, "statusLabel": STATUS_LABEL[OUT_OF_NORMAL],
                "threshold": high, "deviation": round(value - high, 4)}
    return {"status": NORMAL, "statusLabel": STATUS_LABEL[NORMAL]}


def check(store, object_id: str, q: dict) -> dict[str, Any]:
    """取对象的实时值，逐点按限值判断。需要实时值接口（实时历史接口补丁）。"""
    items = _load()
    if not items:
        return {**_empty_payload(), "objectId": object_id}
    try:
        from backend.services import timeseries
    except ImportError:  # pragma: no cover
        return {"objectId": object_id, "available": False,
                "emptyReason": "越限判断需要实时值接口（实时历史接口补丁），当前未安装。"}

    limits = limits_for_object(store, object_id)
    if not limits["items"]:
        return {**limits, "checked": 0, "results": [],
                "notice": "该对象没有可用的工艺限值，无法做越限判断。"}

    realtime = timeseries.realtime(store, {**q, "limit": q.get("limit") or "200"}, object_id)
    by_tag = {item["tagId"]: item for item in realtime["items"]}

    results = []
    counter: dict[str, int] = {}
    for limit in limits["items"]:
        targets, _ = resolve_targets(store, limit)
        for tag in targets:
            reading = by_tag.get(tag["id"])
            if reading is None:
                continue
            verdict = evaluate(limit, reading.get("value"))
            counter[verdict["status"]] = counter.get(verdict["status"], 0) + 1
            results.append({
                "limitId": limit["id"], "appliesTo": limit["appliesTo"], "operatingMode": limit["operatingMode"],
                "limitStatus": limit["status"], "confirmed": limit["confirmed"], "confidence": limit["confidence"],
                "unit": limit["unit"] or tag.get("unit"),
                "range": {k: limit[k] for k in ("lowLow", "low", "normalLow", "normalHigh", "high", "highHigh")},
                "tagId": tag["id"], "tagName": tag["name"], "description": tag.get("rawDescription"),
                "equipmentName": tag.get("equipmentName"),
                "value": reading.get("value"), "quality": reading.get("quality"), "timeText": reading.get("timeText"),
                **verdict,
            })

    return {
        "scope": limits["scope"], "checked": len(results), "summary": counter, "results": results,
        "limitCount": limits["total"], "available": True,
        "dataSource": realtime.get("dataSource"),
        "sourceType": SOURCE_TYPE, "source": {"file": LIMIT_FILE.name, "sheet": SHEET},
        "notice": "限值状态非「已确认」时，越限结论仅供工艺调查参考，不能当作报警。",
    }

def check_limit(store, limit_id: str, q: dict) -> dict[str, Any]:
    """只检查一条限值：取它覆盖的测点实时值，逐点判断。界面展开某一行时用。"""
    items = _load()
    if not items:
        return {**_empty_payload(), "limitId": limit_id}
    for entry in items:
        _resolve_limit_area(store, entry)
    limit = next((i for i in items if i["id"] == limit_id), None)
    if limit is None:
        return {"limitId": limit_id, "available": False, "emptyReason": f"没有编号为 {limit_id} 的限值"}

    targets, notes = resolve_targets(store, limit)
    if not targets:
        return {"limitId": limit_id, "limit": limit, "available": True, "checked": 0, "results": [],
                "notes": notes, "notice": "该限值没有落到任何测点，无法取当前值。"}

    try:
        from backend.services import timeseries
    except ImportError:  # pragma: no cover
        return {"limitId": limit_id, "limit": limit, "available": False,
                "emptyReason": "取当前值需要实时值接口（实时历史接口补丁），当前未安装。"}

    cap = min(len(targets), 200)
    readings = timeseries.realtime(store, {"tagIds": ",".join(t["id"] for t in targets[:cap]), "limit": str(cap)})
    by_tag = {item["tagId"]: item for item in readings["items"]}

    results, counter = [], {}
    for tag in targets[:cap]:
        reading = by_tag.get(tag["id"], {})
        verdict = evaluate(limit, reading.get("value"))
        counter[verdict["status"]] = counter.get(verdict["status"], 0) + 1
        results.append({
            **_target_brief(tag),
            "value": reading.get("value"), "quality": reading.get("quality"),
            "timeText": reading.get("timeText"), **verdict,
        })
    return {
        "limitId": limit_id, "limit": limit, "available": True, "checked": len(results),
        "summary": counter, "results": results, "notes": notes,
        "dataSource": readings.get("dataSource"),
        "notice": "限值状态非「已确认」时，越限结论仅供工艺调查参考，不能当作报警。",
    }
