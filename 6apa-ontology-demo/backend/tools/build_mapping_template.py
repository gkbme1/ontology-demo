"""生成人工对象映射表 6APA_对象映射.xlsx（需求文档 9.3 第三优先级）。

关系 Excel 使用工艺抽象名称，DCS 点表使用现场位号命名，两者零条精确匹配。
本脚本按工段与工艺语义编制首版映射，并在写出前逐行校验目标本体对象
是否真实存在于 DCS 点表中；无法解析的行会被剔除并打印告警，
从而保证映射表里不会出现 DCS 点表中不存在的设备或参数。

所有行的“确认状态”均为“待现场确认”，不得视为已核实的现场事实。
"""
from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from backend.services.dcs_loader import load_dcs  # noqa: E402

SOURCE_DIR = ROOT / "data" / "source"
DCS_FILE = SOURCE_DIR / "6APA_DCS.xlsx"
OUT_FILE = SOURCE_DIR / "6APA_对象映射.xlsx"

COLUMNS = [
    "关系对象语义名称 (Relation Entity)",
    "关系对象类型 (Relation Entity Type)",
    "关系所属工段 (Relation Area)",
    "标准本体对象类型 (Ontology Object Type)",
    "标准本体对象名称 (Ontology Object Name)",
    "本体对象标识 (OWL Subject URI / Area)",
    "本体参数语义名称 (Parameter Name)",
    "置信度 (Confidence)",
    "确认状态 (Status)",
    "映射说明 (Mapping Note)",
]

TANKS = list("ABCDEFGHIJK")
CENTRIFUGES = ["B1", "B2", "B3", "B4", "B5"]
PENDING = "待现场确认"


def rows() -> list[dict]:
    out: list[dict] = []

    def add(entity, etype, area, otype, oname, ouri, pname, conf, note):
        out.append(
            {
                COLUMNS[0]: entity,
                COLUMNS[1]: etype,
                COLUMNS[2]: area,
                COLUMNS[3]: otype,
                COLUMNS[4]: oname,
                COLUMNS[5]: ouri,
                COLUMNS[6]: pname,
                COLUMNS[7]: conf,
                COLUMNS[8]: PENDING,
                COLUMNS[9]: note,
            }
        )

    # ---------------------------------------------------------------- 工段对齐
    add("水解工段", "工段", "水解工段", "工段", "酶法水解工段 (Enzymatic Hydrolysis)",
        "酶法水解工段 (Enzymatic Hydrolysis)", None, "高",
        "关系 Excel 的“水解工段”对应 DCS 点表控制站 13 的酶法水解工段")
    add("萃取工段", "工段", "萃取工段", "工段", "提取水洗工段 (Extraction & Washing)",
        "提取水洗工段 (Extraction & Washing)", None, "高",
        "关系 Excel 的“萃取工段”对应 DCS 点表控制站 10 的提取水洗工段")

    # ---------------------------------------------------------------- 生产设备
    for t in TANKS:
        add("水解反应釜", "生产设备", "水解工段", "设备", f"水解罐{t}", f":Dev_13_水解罐{t}", None, "高",
            f"酶法水解工段共 11 台并联酶法水解罐（A~K），关系表中的“水解反应釜”为其工艺抽象类；本行对应第 {t} 台")
    add("水解田缓冲罐", "生产设备", "水解工段", "设备", "水解液缓冲罐", ":Dev_13_水解液缓冲罐", None, "中",
        "关系表所述“水解田缓冲罐”即水解终点出料暂存罐，DCS 命名为“水解液缓冲罐”")
    for c in CENTRIFUGES:
        add("水洗碟片离心机组", "生产设备", "萃取工段", "设备", f"离心机{c}", f":Dev_10_离心机{c}", None, "高",
            f"提取水洗工段 5 台并联碟片离心机 B1~B5，关系表按机组抽象；本行对应 {c}")
    add("进料预冷换热器", "生产设备", "萃取工段", "设备", "水解液缓冲罐", ":Dev_13_水解液缓冲罐", None, "中",
        "现场该换热器自身无独立位号，其出口温度与冷盐水调节阀测点均挂在“水解液缓冲罐”设备下")
    add("进料预冷换热器", "生产设备", "萃取工段", "设备", "丁酯换热器", ":Dev_10_丁酯换热器", None, "低",
        "萃取侧丁酯换热器，是否属于关系表所指“进料预冷换热器”需现场确认")
    add("6-APA结晶反应釜", "生产设备", "跨工段协同", "设备", "APA结晶罐", ":Dev_14_APA结晶罐", None, "中",
        "下游结晶工段主结晶单元，DCS 命名为“APA结晶罐”")
    add("苯乙酸精馏回收塔", "生产设备", "跨工段协同", "设备", "丁酯精馏塔", ":Dev_17_丁酯精馏塔", None, "低",
        "苯乙酸回收精馏工段内的精馏塔主体，现场 DCS 命名为“丁酯精馏塔”，是否为关系表所指回收塔需现场确认")

    # ---------------------------------------------------------------- 设备部件
    for t in TANKS:
        add("机械搅拌装置", "设备部件", "水解工段", "参数", f"水解罐{t}", f":Dev_13_水解罐{t}", "搅拌", "中",
            "现场该部件无独立设备位号，仅以水解罐搅拌电机控制点（MOT）体现")

    # ---------------------------------------------------------------- 工艺属性
    for t in TANKS:
        add("釜内反应温度", "工艺属性", "水解工段", "参数", f"水解罐{t}", f":Dev_13_水解罐{t}", "温度", "高",
            "水解罐温度检测点，工程单位 ℃")
        for n in (1, 2):
            add("反应体系pH", "工艺属性", "水解工段", "参数", f"水解罐{t}{n}", f":Dev_13_水解罐{t}{n}", "pH", "高",
                f"水解罐{t}第 {n} 路 pH 检测点，工程单位 ph")
    add("RB进料流量", "工艺属性", "水解工段", "参数", "RB进水解罐", ":Dev_13_RB进水解罐", "流量A", "高",
        "底物 RB 浓缩液进水解罐流量 A 路")
    add("RB进料流量", "工艺属性", "水解工段", "参数", "RB进水解罐", ":Dev_13_RB进水解罐", "流量B", "高",
        "底物 RB 浓缩液进水解罐流量 B 路")
    add("硼酸进料流量", "工艺属性", "水解工段", "参数", "公用测量管线", ":Dev_13_公用测量管线", "硼酸流量A", "高",
        "硼酸缓冲液进料流量 A 路")
    add("硼酸进料流量", "工艺属性", "水解工段", "参数", "公用测量管线", ":Dev_13_公用测量管线", "硼酸流量B", "高",
        "硼酸缓冲液进料流量 B 路")
    add("氨水加料流量", "工艺属性", "水解工段", "参数", "公用测量管线", ":Dev_13_公用测量管线", "氨水流量计", "高",
        "pH 中和滴加氨水流量计")
    add("热水罐出水温度", "工艺属性", "水解工段", "参数", "热水罐", ":Dev_13_热水罐", "温度", "高",
        "水解工段热水罐温度检测点")
    add("热水罐出水温度", "工艺属性", "水解工段", "参数", "热水罐B", ":Dev_13_热水罐B", "温度检测", "高",
        "水解工段热水罐 B 温度检测点")
    for s in ("a", "b", "c"):
        add("水解进料换热温度", "工艺属性", "萃取工段", "参数", "水解液缓冲罐", ":Dev_13_水解液缓冲罐",
            f"去萃取换热器出口温度{s}", "高", f"水解液去萃取换热器出口温度 {s} 路，即萃取进料换热后温度")
    for s in ("a", "b"):
        add("换热器冷盐水流量", "工艺属性", "萃取工段", "参数", "水解液缓冲罐", ":Dev_13_水解液缓冲罐",
            f"去萃取换热器冷盐水调节阀{s}反馈", "低",
            "替代测点：现场无冷盐水流量计，暂以去萃取换热器冷盐水调节阀开度反馈（%）作为间接测量锚点")
    for s in ("A", "B"):
        add("丁酯进料流量", "工艺属性", "萃取工段", "参数", "公用测量管线", ":Dev_10_公用测量管线",
            f"丁酯进料流量{s}", "高", f"萃取剂醋酸丁酯进料流量 {s} 路")
    add("破乳剂加入量", "工艺属性", "萃取工段", "参数", "公用测量管线", ":Dev_10_公用测量管线",
        "破乳剂备用流量", "中", "破乳剂总管流量测点")
    for s in ("（1—2）", "（3—4）"):
        add("破乳剂加入量", "工艺属性", "萃取工段", "参数", "离心机200", ":Dev_10_离心机200",
            f"{s}破乳剂进料流量", "中", f"离心机200 {s} 机组破乳剂进料流量")
    for c in CENTRIFUGES:
        add("轻相出口背压", "工艺属性", "萃取工段", "参数", f"离心机{c}", f":Dev_10_离心机{c}",
            "轻相出口压力", "高", f"离心机{c}轻相（有机相）出口压力，工程单位 Mpa")
    for s in "ABCDE":
        add("稀硫酸进料流量", "工艺属性", "萃取工段", "参数", "公用测量管线", ":Dev_12_公用测量管线",
            f"混合器M2123{s}稀硫酸流量(DN25)", "中",
            f"酸化用稀硫酸流量测点位于配制与薄膜蒸发工段混合器 M2123{s}，跨工段引用")

    # ---------------------------------------------------------------- 执行机构
    for t in TANKS:
        add("氨水加料气动调节阀", "执行机构", "水解工段", "参数", f"{t}罐", f":Dev_13_{t}罐", "氨水进料阀", "中",
            f"{t} 罐氨水进料阀（VAL 块），承担 pH 中和氨水滴加执行")
        add("水解罐进出料气动切断阀", "执行机构", "水解工段", "参数", f"{t}罐RB", f":Dev_13_{t}罐RB", "进料阀", "中",
            f"{t} 罐 RB 进料切断阀（VAL 块）")
    add("夹套伴热蒸汽/冷水调节阀", "执行机构", "水解工段", "参数", "热水罐B", ":Dev_13_热水罐B", "温度控制阀", "低",
        "现场未见水解罐夹套调节阀独立位号，暂以热水罐 B 温度控制阀作为伴热侧执行机构候选")
    for s in ("a", "b"):
        add("换热器冷盐水调节阀", "执行机构", "萃取工段", "参数", "水解液缓冲罐", ":Dev_13_水解液缓冲罐",
            f"去萃取换热器冷盐水调节阀{s}指令", "中", f"去萃取换热器冷盐水调节阀 {s} 路指令（AO 块）")
    for c in CENTRIFUGES:
        add("丁酯溶媒进料调节阀", "执行机构", "萃取工段", "参数", f"离心机{c}", f":Dev_10_离心机{c}",
            "丁酯进口调节阀指令", "中", f"离心机{c}丁酯进口调节阀指令（AO 块）")

    return out


def main() -> int:
    ds = load_dcs(DCS_FILE)
    from backend.ontology import identifiers as ids

    kept, dropped = [], []
    for row in rows():
        otype = row[COLUMNS[3]]
        uri = row[COLUMNS[5]]
        pname = row[COLUMNS[6]]
        if otype == "工段":
            ok = any(a["name"] == uri for a in ds.areas.values())
        elif otype == "设备":
            ok = ds.equipment_by_owl(uri) is not None
        elif otype == "参数":
            ok = ids.parameter_id(uri, pname) in ds.parameters
        else:
            ok = False
        (kept if ok else dropped).append(row)

    if dropped:
        print(f"⚠ 以下 {len(dropped)} 行在 DCS 点表中无法解析，已剔除：")
        for row in dropped:
            print("   -", row[COLUMNS[0]], "→", row[COLUMNS[5]], row[COLUMNS[6]] or "")

    df = pd.DataFrame(kept, columns=COLUMNS)
    with pd.ExcelWriter(OUT_FILE, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="对象映射", index=False)
        ws = writer.sheets["对象映射"]
        fill = PatternFill("solid", fgColor="1F3A5F")
        for j, _ in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=1, column=j)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(j)].width = [26, 18, 16, 20, 22, 30, 28, 12, 14, 60][j - 1]
        ws.freeze_panes = "A2"

    print(f"✓ 写出 {OUT_FILE.name}：{len(kept)} 行有效映射（全部标记为“{PENDING}”）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
