"""生成《6APA_工艺限值.xlsx》与配套空模板。

两点边界必须说清楚：
1. 本表**不是**现场工艺规程。带数值的行是"推断初版"，依据写在「限值依据」列，
   全部标记「待现场确认」，需工艺人员逐条核对后才能作为正式依据。
2. 没有可靠推断依据的行**故意留空数值**，只保留对象、工况与单位，
   作为工艺调查的待填清单——留空比填一个编出来的数字有用得多。

运行：python3 -m backend.tools.build_limit_table
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "data" / "source" / "6APA_工艺限值.xlsx"
TEMPLATE = ROOT / "data" / "source" / "6APA_工艺限值.模板.xlsx"

COLUMNS = [
    ("限值编号 (Limit ID)", 14),
    ("适用对象类型 (Applies To Type)", 18),
    ("适用对象 (Applies To)", 22),
    ("适用工段 (Area)", 14),
    ("适用工况 (Operating Mode)", 16),
    ("参数含义 (Parameter)", 16),
    ("单位 (Unit)", 10),
    ("低低限 LL", 10),
    ("低限 L", 10),
    ("正常下限 (Normal Low)", 14),
    ("正常上限 (Normal High)", 14),
    ("高限 H", 10),
    ("高高限 HH", 10),
    ("限值依据 (Basis)", 52),
    ("依据类型 (Basis Type)", 16),
    ("置信度 (Confidence)", 12),
    ("确认状态 (Status)", 14),
    ("备注 (Note)", 30),
]

PENDING = "待现场确认"
TO_FILL = "待现场提供"

# 每行：对象类型, 对象, 工段, 工况, 参数含义, 单位, LL, L, 正常下, 正常上, H, HH, 依据, 依据类型, 置信度, 状态, 备注
ROWS: list[tuple] = [
    # ---------- 有推断依据的（数值来自关系表工艺描述 + 酶法常规工况，全部待确认）----------
    ("工艺属性", "釜内反应温度", "水解工段", "水解反应期", "温度", "℃",
     30.0, 32.0, 33.0, 35.0, 35.5, 37.0,
     "关系表记载「反应超温(>35℃)导致酰化酶蛋白热变性失活」「超温加速6-APA母环开环水解」，据此取 35℃ 为高限；"
     "正常区间按青霉素G酰化酶常规操作温度 33~35℃ 推断",
     "关系表工艺描述 + 酶法常规", "中", PENDING, "11 台并联水解罐共用，如各罐工艺不同需拆开填写"),

    ("工艺属性", "反应体系pH", "水解工段", "水解反应期", "pH", "pH",
     7.0, 7.5, 7.8, 8.2, 8.4, 8.8,
     "青霉素G酰化酶水解最适 pH 常规在 7.5~8.5，偏酸降低酶活、偏碱加剧6-APA降解；关系表记载氨水按 pH 动态微量滴加",
     "酶法常规 + 关系表工艺描述", "中", PENDING, "每罐 2 支 pH 计（A1/A2 等），双支偏差阈值需另行确认"),

    ("工艺属性", "水解进料换热温度", "萃取工段", "进萃取前", "温度", "℃",
     2.0, 4.0, 6.0, 12.0, 15.0, 18.0,
     "关系表记载「萃取来料超温加剧有机溶剂挥发与相界面乳化」，据此设上限；下限按冷盐水系统可达温度推断",
     "关系表工艺描述 + 推断", "低", PENDING, "具体降温目标需按萃取工艺卡确认"),

    ("DCS点", "TI_1407*", "结晶工段", "结晶期", "温度", "℃",
     0.0, 3.0, 5.0, 12.0, 15.0, 18.0,
     "6-APA 低温结晶常规温度区间；本体中该位号族为 14 台结晶罐温度检测",
     "工艺常规推断", "低", PENDING, "结晶温度直接影响收率与晶型，务必按现场结晶曲线确认"),

    ("DCS点", "AT_1401*", "结晶工段", "结晶期", "pH", "pH",
     3.6, 3.9, 4.0, 4.3, 4.5, 4.8,
     "6-APA 等电点约 pH 4.3，等电点结晶常规控制在 4.0~4.3；本体中该位号族为 14 台结晶罐 pH 检测",
     "工艺常规推断", "低", PENDING, "等电点随温度与离子强度变化，需现场确认"),

    ("工艺属性", "酸化体系pH", "萃取工段", "酸化萃取期", "pH", "pH",
     1.2, 1.8, 2.0, 2.5, 2.8, 3.2,
     "酸化萃取常规将水相调至强酸区以利苯乙酸转入有机相；该对象在 DCS 点表中未匹配到测点",
     "工艺常规推断", "低", PENDING, "无对应测点，需确认是否有在线 pH 或仅离线分析"),

    ("工艺属性", "轻相出口背压", "萃取工段", "离心机运行期", "压力", "MPa",
     0.05, 0.10, 0.15, 0.30, 0.35, 0.45,
     "按离心机出口背压常规操作区间推断；本体已匹配 5 台离心机的轻相出口压力测点",
     "设备常规推断", "低", PENDING, "背压直接影响分相界面位置，应按离心机厂家参数确认"),

    ("工艺属性", "热水罐出水温度", "水解工段", "常态", "温度", "℃",
     40.0, 45.0, 50.0, 60.0, 65.0, 70.0,
     "热水系统常规供水温度区间推断；本体已匹配热水罐、热水罐B 两支温度计",
     "推断", "低", PENDING, "用途（伴热/清洗/配液）不同则区间不同，需确认"),

    # ---------- 无可靠依据、留空数值，作为工艺调查待填清单 ----------
    ("工艺属性", "釜内反应液位", "水解工段", "水解反应期", "液位", "", None, None, None, None, None, None,
     "关系表中存在该工艺属性，但未匹配到 DCS 测点，且装料量随批量变化", "待现场提供", "", TO_FILL,
     "请现场确认正常装料液位区间与联锁高低限"),
    ("工艺属性", "搅拌运行转速", "水解工段", "水解反应期", "转速", "", None, None, None, None, None, None,
     "关系表记载搅拌转速影响固液传质速率，但未匹配到转速测点", "待现场提供", "", TO_FILL, ""),
    ("工艺属性", "RB进料流量", "水解工段", "进料过程", "流量", "L/h", None, None, None, None, None, None,
     "流量设定随批量与投料方案变化，无法推断", "待现场提供", "", TO_FILL, "已匹配 FT_1302A/B 两支流量计"),
    ("工艺属性", "硼酸进料流量", "水解工段", "进料过程", "流量", "L/h", None, None, None, None, None, None,
     "缓冲液配比属工艺配方，需现场提供", "待现场提供", "", TO_FILL, "已匹配 FT_1301A/B"),
    ("工艺属性", "氨水加料流量", "水解工段", "中和过程", "流量", "m3/h", None, None, None, None, None, None,
     "按 pH 闭环动态调节，限值需现场确认", "待现场提供", "", TO_FILL, "已匹配 FT_1315"),
    ("工艺属性", "丁酯进料流量", "萃取工段", "进料过程", "流量", "m3/h", None, None, None, None, None, None,
     "溶媒比随批量变化，需现场提供", "待现场提供", "", TO_FILL, "已匹配 FT_DZ_A/B"),
    ("工艺属性", "破乳剂加入量", "萃取工段", "离心分离", "流量", "m3/h", None, None, None, None, None, None,
     "加药量属工艺配方，需现场提供", "待现场提供", "", TO_FILL, "已匹配 FT_1101A/B、FT_1102"),
    ("工艺属性", "稀硫酸进料流量", "萃取工段", "加酸过程", "流量", "m3/h", None, None, None, None, None, None,
     "加酸量按 pH 目标调节，需现场提供", "待现场提供", "", TO_FILL, "已匹配 FT_1106A~E"),
    ("工艺属性", "换热器冷盐水流量", "萃取工段", "换热过程", "阀位/流量", "%", None, None, None, None, None, None,
     "本体中该对象锚定的是调节阀开度反馈，非流量计，量纲需先确认", "待现场提供", "", TO_FILL,
     "已匹配 HV_1016A/B_FB（开度 %）"),
    ("工艺属性", "离心机转速", "萃取工段", "离心分离", "转速", "rpm", None, None, None, None, None, None,
     "按离心机型号确定，未匹配到转速测点", "待现场提供", "", TO_FILL, ""),
    ("工艺属性", "重相出口背压", "萃取工段", "离心机运行期", "压力", "MPa", None, None, None, None, None, None,
     "与轻相背压配合确定分相界面，需现场成对确认", "待现场提供", "", TO_FILL, "未匹配到测点"),
    ("质量指标", "水解酶解转化率", "水解工段", "批次结束", "转化率", "%", None, None, None, None, None, None,
     "属批次质量指标，目标值由工艺部门下达", "待现场提供", "", TO_FILL, "离线分析，无在线测点"),
    ("质量指标", "废有机相6-APA夹带残效", "萃取工段", "批次结束", "残留", "%", None, None, None, None, None, None,
     "属批次质量指标，需现场给出可接受上限", "待现场提供", "", TO_FILL, "离线分析"),
    ("业务指标", "产品批次提取收率", "萃取工段", "批次结束", "收率", "%", None, None, None, None, None, None,
     "属考核指标，目标值由生产管理部门下达", "待现场提供", "", TO_FILL, ""),
]


def _write(path: Path, rows: list[tuple]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "工艺限值"
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DCE6F1")
    for index, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=index, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "A2"

    for offset, row in enumerate(rows):
        values = (f"PL{offset + 1:03d}",) + row
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=offset + 2, column=col, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=col in (14, 18))

    notes = wb.create_sheet("字段说明")
    notes.append(["字段", "是否必填", "填写说明"])
    for name, required, desc in [
        ("限值编号 (Limit ID)", "必填", "唯一编号，例如 PL001"),
        ("适用对象类型 (Applies To Type)", "必填", "工艺属性 / 质量指标 / 业务指标 / 设备 / 参数 / DCS点"),
        ("适用对象 (Applies To)", "必填",
         "对象名称。工艺属性填关系表里的名称（如 釜内反应温度）；DCS点可填位号，支持 * 通配，如 TI_1407*"),
        ("适用工段 (Area)", "选填", "如 水解工段、萃取工段、结晶工段"),
        ("适用工况 (Operating Mode)", "选填", "如 水解反应期、进料过程、待机；同一对象不同工况可写多行"),
        ("参数含义 (Parameter)", "选填", "温度 / pH / 压力 / 流量 / 液位 等"),
        ("单位 (Unit)", "选填", "℃、pH、MPa、m3/h 等；留空则取测点单位"),
        ("低低限 LL / 低限 L / 高限 H / 高高限 HH", "选填", "四档报警限；留空表示该档未设"),
        ("正常下限 / 正常上限", "选填", "正常操作区间"),
        ("限值依据 (Basis)", "必填", "这个数字从哪来。现场填写时请写工艺规程名称与条款号"),
        ("依据类型 (Basis Type)", "选填", "工艺规程 / 关系表工艺描述 / 设备铭牌 / 推断 / 待现场提供"),
        ("置信度 (Confidence)", "选填", "高 / 中 / 低；来自工艺规程的填 高"),
        ("确认状态 (Status)", "必填", "待现场确认 / 待现场提供 / 已确认；只有「已确认」才能作为正式依据"),
        ("备注 (Note)", "选填", "补充说明"),
    ]:
        notes.append([name, required, desc])
    notes.column_dimensions["A"].width = 34
    notes.column_dimensions["B"].width = 12
    notes.column_dimensions["C"].width = 72

    usage = wb.create_sheet("使用说明")
    usage.append(["项目", "说明"])
    for k, v in [
        ("用途", "录入 6-APA 生产线的工艺限值（上下限 / 四档报警限），供本体接口、界面与工艺调查使用。"),
        ("生效方式", "文件放在 data/source/6APA_工艺限值.xlsx，重启本体服务即自动加载，无需改代码。"),
        ("绑定方式", "限值挂在工艺属性上，通过本体已有的测量锚点自动落到具体 DCS 位号；也可直接按位号（支持 * 通配）写特例。"),
        ("数据边界", "本体不生成限值。带数值的行是推断初版，标记「待现场确认」；留空数值的行是待填清单，"
                    "标记「待现场提供」。现场核对后把状态改为「已确认」并补上依据出处。"),
        ("注意", "限值一旦标为「已确认」，就会被当作正式工艺依据参与越限判断，请谨慎。"),
    ]:
        usage.append([k, v])
    usage.column_dimensions["A"].width = 14
    usage.column_dimensions["C" if False else "B"].width = 96

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    _write(OUT, ROWS)
    _write(TEMPLATE, [])
    filled = sum(1 for r in ROWS if r[9] is not None or r[10] is not None)
    print(f"已生成 {OUT}：{len(ROWS)} 行（{filled} 行带推断数值，{len(ROWS) - filled} 行待现场提供）")
    print(f"已生成 {TEMPLATE}：空模板")


if __name__ == "__main__":
    main()
