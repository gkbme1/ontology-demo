"""按规则生成设备归并表 data/source/6APA_设备归并.xlsx。

    python3 -m backend.tools.build_equipment_merge          # 生成（已存在时不覆盖）
    python3 -m backend.tools.build_equipment_merge --force  # 按规则重新生成，覆盖人工修改

DCS 点表把同一台现场设备写成了好几种名字（水解罐A / A罐 / A罐RB / 水解罐A1 …），
本脚本在同一标准设备组内按规则找出这些写法，写成一张可编辑的归并表；
本体服务加载 DCS 点表时按表把它们合并成一台设备、统一名称。DCS Excel 本身不改。
规则见 backend/services/equipment_merge.py。表里「启用」改成 否 即不归并，也可以手工追加行。
"""
from __future__ import annotations

import sys
from collections import Counter
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from backend.services import equipment_merge as em  # noqa: E402
from backend.services.dcs_loader import load_dcs  # noqa: E402

SOURCE_DIR = ROOT / "data" / "source"
OUT_FILE = SOURCE_DIR / em.MERGE_FILE_NAME

NOTES = [
    ("用途", "DCS 点表把同一台设备写成了多种名字，本表把它们归并成一台。本体服务加载 DCS 点表时按本表合并，DCS Excel 本身不改。"),
    ("统一设备名", "合并后的设备名称。同组声明了实例名（设备主数据「实例」列）时取实例名，否则取点位最多的写法。"),
    ("统一OWL Subject", "合并后保留的设备标识。原设备的 ID、OWL Subject、名称都登记为别名，旧 ID / 旧名称仍能查到合并后的设备。"),
    ("归并方式", "同一设备：同一台设备的另一种写法；附属点位：设备上的测量元件或信号（如两支 pH 电极、RB 进料阀、故障信号）；统一名称：只改名不合并。"),
    ("参数前缀", "附属点位的参数名前加的前缀，避免与本体参数重名，例如 A罐RB 的「进料阀指令」→「RB进料阀指令」。同名参数会合并成一个参数、保留全部位号。"),
    ("启用", "是 / 否。改成 否 即不归并该行；也可以照格式手工追加行。"),
    ("重新生成", "python3 -m backend.tools.build_equipment_merge --force 会按规则覆盖本表，人工修改会丢失。"),
    ("规则", "本名 / 同名异写 / 简称 / 冗余后缀 / 省略设备位号 / 附属点位 / 同名跨站，均限定在同一标准设备组内，详见 backend/services/equipment_merge.py。"),
]


def main() -> int:
    force = "--force" in sys.argv
    if OUT_FILE.exists() and not force:
        print(f"{OUT_FILE.name} 已存在，未覆盖（按规则重建请加 --force）")
        return 0
    ds = load_dcs(SOURCE_DIR / "6APA_DCS.xlsx")  # 不带归并表，拿原始写法
    rows = em.propose(ds)
    df = pd.DataFrame(rows, columns=em.COLUMNS)
    with pd.ExcelWriter(OUT_FILE, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=em.SHEET, index=False)
        pd.DataFrame(NOTES, columns=["项目", "说明"]).to_excel(writer, sheet_name="说明", index=False)
        ws = writer.sheets[em.SHEET]
        widths = [6, 7, 26, 20, 8, 28, 8, 20, 28, 10, 12, 60, 10, 6]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        head = PatternFill("solid", fgColor="1F3A5F")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head
            cell.alignment = Alignment(vertical="center")
        ws.freeze_panes = "D2"
        ws.auto_filter.ref = ws.dimensions
        notes = writer.sheets["说明"]
        notes.column_dimensions["A"].width = 16
        notes.column_dimensions["B"].width = 110
    kinds = Counter(r["归并方式"] for r in rows)
    merged = sum(1 for r in rows if r["归并方式"] != em.KIND_RENAME)
    print(f"已生成 {OUT_FILE.relative_to(ROOT)}：{len(rows)} 行（{dict(kinds)}），"
          f"设备 {len(ds.equipment)} → {len(ds.equipment) - merged}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
